#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generate adoption market share data from:
1. 학교별 담당총판 및 본사담당자 정보.xlsx  - student counts per school
2. 21-26주문데이터.xlsx - order data

Output: public/public/assets/adoption_data.json
"""
import openpyxl
from collections import defaultdict
import json
import os

BASE = os.path.dirname(os.path.abspath(__file__))
SCHOOL_FILE = os.path.join(BASE, 'public', 'public', 'assets', '학교별 담당총판 및 본사담당자 정보.xlsx')
ORDER_FILE  = os.path.join(BASE, 'public', 'public', 'assets', '21-26주문데이터.xlsx')
OUT_FILE    = os.path.join(BASE, 'public', 'public', 'assets', 'adoption_data.json')

SUDOGWON = {'서울특별시', '경기도', '인천광역시'}

# ------------------------------------------------------------------
# 1. Load school student counts
# ------------------------------------------------------------------
print("Loading school info...")
wb_school = openpyxl.load_workbook(SCHOOL_FILE, read_only=True)
ws_school = wb_school['01.학교정보']

# Key: sido_name (시도교육청 based on first column)
# Value: {g1, g2, g3}
sido_students = defaultdict(lambda: {'g1': 0.0, 'g2': 0.0, 'g3': 0.0})

# Map 시도교육청 -> 시도명 (for matching with order data)
sido_edu_to_sido = {
    '서울특별시교육청': '서울특별시',
    '부산광역시교육청': '부산광역시',
    '대구광역시교육청': '대구광역시',
    '인천광역시교육청': '인천광역시',
    '광주광역시교육청': '광주광역시',
    '대전광역시교육청': '대전광역시',
    '울산광역시교육청': '울산광역시',
    '세종특별자치시교육청': '세종특별자치시',
    '경기도교육청': '경기도',
    '강원특별자치도교육청': '강원특별자치도',
    '충청북도교육청': '충청북도',
    '충청남도교육청': '충청남도',
    '전북특별자치도교육청': '전북특별자치도',
    '전라남도교육청': '전라남도',
    '경상북도교육청': '경상북도',
    '경상남도교육청': '경상남도',
    '제주특별자치도교육청': '제주특별자치도',
}

for i, row in enumerate(ws_school.iter_rows(min_row=2, values_only=True)):
    edu_sido = row[0]  # 시도교육청
    sido = sido_edu_to_sido.get(edu_sido, edu_sido)
    if sido is None:
        continue
    g1 = float(row[10] or 0)
    g2 = float(row[13] or 0)
    g3 = float(row[16] or 0)
    sido_students[sido]['g1'] += g1
    sido_students[sido]['g2'] += g2
    sido_students[sido]['g3'] += g3

total_students_all = {
    'g1': sum(v['g1'] for v in sido_students.values()),
    'g2': sum(v['g2'] for v in sido_students.values()),
    'g3': sum(v['g3'] for v in sido_students.values()),
}
total_students_all['total'] = total_students_all['g1'] + total_students_all['g2'] + total_students_all['g3']
# 1개 코호트(학년 평균) = 전체 / 3 → 점유율 분모로 사용
total_students_all['cohort'] = total_students_all['total'] / 3.0

sudogwon_students = {
    'g1': sum(sido_students[s]['g1'] for s in SUDOGWON if s in sido_students),
    'g2': sum(sido_students[s]['g2'] for s in SUDOGWON if s in sido_students),
    'g3': sum(sido_students[s]['g3'] for s in SUDOGWON if s in sido_students),
}
sudogwon_students['total'] = sudogwon_students['g1'] + sudogwon_students['g2'] + sudogwon_students['g3']
sudogwon_students['cohort'] = sudogwon_students['total'] / 3.0

print(f"Total nationwide students: G1={total_students_all['g1']:,.0f}, G2={total_students_all['g2']:,.0f}, G3={total_students_all['g3']:,.0f}")
print(f"Sudogwon students: G1={sudogwon_students['g1']:,.0f}, G2={sudogwon_students['g2']:,.0f}")

# ------------------------------------------------------------------
# 1b. Load region(구) student counts from sales_staff.csv (고등 + 중학교 분리)
# ------------------------------------------------------------------
import csv
STAFF_CSV = os.path.join(BASE, 'public', 'public', 'sales_staff.csv')
region_students    = defaultdict(float)  # region -> 고등학생 코호트
ms_region_students = defaultdict(float)  # region -> 중학교 학생 코호트
try:
    with open(STAFF_CSV, encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        for row_s in reader:
            grade_code = str(row_s.get('학교급코드', '')).strip()
            region = str(row_s.get('지역', '')).strip()
            if not region:
                continue
            try:
                g1 = float(row_s.get('1학년 학생수') or 0)
                g2 = float(row_s.get('2학년 학생수') or 0)
                g3 = float(row_s.get('3학년 학생수') or 0)
                cohort = (g1 + g2 + g3) / 3.0
                if grade_code == '4':   # 고등학교
                    region_students[region] += cohort
                elif grade_code == '3': # 중학교
                    ms_region_students[region] += cohort
            except:
                pass
    print(f"Region student counts loaded: HS={len(region_students)} regions, MS={len(ms_region_students)} regions")
except Exception as e:
    print(f"Warning: could not load region students from CSV: {e}")

# ------------------------------------------------------------------
# 1c. Load staff -> regions mapping + school meta per region
# ------------------------------------------------------------------
STAFF_MAPPING = os.path.join(BASE, 'public', 'public', 'sales_staff_mapping.json')
staff_regions_map = {}  # staff_code -> [region, ...]
# region -> [ {name, code, students, type} ]
region_school_meta = defaultdict(list)
try:
    with open(STAFF_MAPPING, encoding='utf-8') as f:
        sm = json.load(f)
    for code, regions_dict in sm.items():
        staff_regions_map[code] = list(regions_dict.keys())
        for region, schools in regions_dict.items():
            for s in schools:
                meta = s.get('meta', {})
                school_type = meta.get('학교특성', '') or ''
                # 학교특성이 있는 경우만 고등학교(type에 '고등' 포함하거나 비어있지 않은 경우)
                # 중학교/초등학교는 학교특성이 빈 문자열임
                if not school_type:
                    continue
                stu = int(meta.get('학생수계', 0) or 0)
                region_school_meta[region].append({
                    'name': s.get('school', ''),
                    'code': s.get('code', ''),
                    'students': stu,
                    'type': school_type,
                    'founded': meta.get('설립구분', ''),
                })
    print(f"Staff regions loaded: {len(staff_regions_map)} staff, {len(region_school_meta)} regions with school meta")
except Exception as e:
    print(f"Warning: could not load staff mapping: {e}")

# ------------------------------------------------------------------
# 2. Load 2026 textbook order data (고등 + 중학교)
# ------------------------------------------------------------------
print("\nLoading order data (2026 고등+중학 교과서)...")
wb_order = openpyxl.load_workbook(ORDER_FILE, read_only=True)
ws_order = wb_order['06. 21-26주문']

# 중학교 특정 가준 과목 (불필요 과목 제외하여 JSON 크기 절약)
MS_TARGET = {
    '정보', '데이터 분석과 인공지능', '앱과 코딩',
    '디지털 리터러시', '슬기로운 디지털 생활', '슬기로운 인공지능 윤리생활',
    '진로와 직업', '보건', '체육 ①', '체육 ②',
}
HS_INFO_TARGET = {
    '정보', '인공지능 기초', '데이터 과학', '소프트웨어와 생활', '정보과학',
}

# 고등 Aggregates
subject_by_sido   = defaultdict(lambda: defaultdict(int))
subject_by_region = defaultdict(lambda: defaultdict(int))
subject_by_school = defaultdict(lambda: defaultdict(int))
subject_jibang    = defaultdict(int)
subject_total     = defaultdict(int)
subject_sudogwon  = defaultdict(int)
kyogakgun_map     = {}  # subject -> 교과군
sido_covered      = set()
region_covered    = set()

# 중학교 Aggregates
ms_subject_by_region = defaultdict(lambda: defaultdict(int))
ms_subject_jibang    = defaultdict(int)
ms_subject_total     = defaultdict(int)
ms_subject_sudogwon  = defaultdict(int)
ms_kyogakgun_map     = {}
ms_region_covered    = set()

for row in ws_order.iter_rows(min_row=2, values_only=True):
    year = row[0]
    if year != 2026:
        continue
    school_grade = str(row[17] or '')
    book_type    = str(row[19] or '')
    if '교과서' not in book_type:
        continue
    curriculum = str(row[18] or '').strip()
    if curriculum != '2022':
        continue
    sido       = row[3]
    if sido is None:
        continue
    region     = str(row[5] or '').strip()
    subject    = (row[21] or '기타').strip()
    qty        = int(row[23] or 0)
    kyogakgun  = (row[26] or '기타').strip()
    school_code = str(row[14] or '').strip()

    if '중학교' in school_grade:
        if subject not in MS_TARGET:
            continue  # 필요 과목만 저장
        ms_subject_total[subject] += qty
        ms_kyogakgun_map[subject] = kyogakgun
        if region:
            ms_subject_by_region[region][subject] += qty
            ms_region_covered.add(region)
        if sido in SUDOGWON:
            ms_subject_sudogwon[subject] += qty
        else:
            ms_subject_jibang[subject] += qty

    elif '고등' in school_grade:
        subject_by_sido[sido][subject] += qty
        subject_total[subject] += qty
        kyogakgun_map[subject] = kyogakgun
        sido_covered.add(sido)
        if region:
            subject_by_region[region][subject] += qty
            region_covered.add(region)
        if sido in SUDOGWON:
            subject_sudogwon[subject] += qty
        else:
            subject_jibang[subject] += qty
        if school_code:
            subject_by_school[school_code][subject] += qty

print(f"Sido covered: {sorted(sido_covered)}")
print(f"HS Regions covered: {len(region_covered)}, MS Regions covered: {len(ms_region_covered)}")
print(f"Top HS subjects by qty:")
for subj, qty in sorted(subject_total.items(), key=lambda x: -x[1])[:10]:
    kg = kyogakgun_map.get(subj, '')
    print(f"  [{kg}] {subj}: {qty:,}")
print(f"Top MS subjects by qty:")
for subj, qty in sorted(ms_subject_total.items(), key=lambda x: -x[1]):
    kg = ms_kyogakgun_map.get(subj, '')
    print(f"  [{kg}] {subj}: {qty:,}")

# ------------------------------------------------------------------
# 3. Compute market shares
# ------------------------------------------------------------------
print("\nComputing market shares...")

def safe_pct(num, denom):
    if denom and denom > 0:
        return round(num / denom * 100, 2)
    return 0.0

# Use 1 cohort (total/3) as denominator: the average number of students per grade
# This gives meaningful percentages for per-subject textbook adoption
total_all_students = total_students_all['cohort']
total_sudo_students = sudogwon_students['cohort']

# Nationwide subject share
nationwide_shares = []
for subj, qty in sorted(subject_total.items(), key=lambda x: -x[1]):
    kg = kyogakgun_map.get(subj, '기타')
    pct = safe_pct(qty, total_all_students)
    nationwide_shares.append({
        'subject': subj,
        'kyogakgun': kg,
        'qty': qty,
        'pct': pct,
    })

# Sudogwon subject share
sudogwon_shares = []
for subj, qty in sorted(subject_sudogwon.items(), key=lambda x: -x[1]):
    kg = kyogakgun_map.get(subj, '기타')
    pct = safe_pct(qty, total_sudo_students)
    sudogwon_shares.append({
        'subject': subj,
        'kyogakgun': kg,
        'qty': qty,
        'pct': pct,
    })

# Regional (sido) subject shares
regional = []
for sido in sorted(sido_covered):
    sido_g1 = sido_students.get(sido, {}).get('g1', 0)
    sido_g2 = sido_students.get(sido, {}).get('g2', 0)
    sido_g3 = sido_students.get(sido, {}).get('g3', 0)
    sido_total = sido_g1 + sido_g2 + sido_g3
    sido_cohort = sido_total / 3.0
    subjects_in_sido = sorted(subject_by_sido[sido].items(), key=lambda x: -x[1])
    subj_list = []
    for subj, qty in subjects_in_sido:
        kg = kyogakgun_map.get(subj, '기타')
        pct = safe_pct(qty, sido_cohort)
        subj_list.append({'subject': subj, 'kyogakgun': kg, 'qty': qty, 'pct': pct})
    regional.append({
        'sido': sido,
        'total_students': sido_total,
        'cohort_students': sido_cohort,
        'subjects': subj_list,
    })

# Build kyogakgun summary (교과군별 집계)
kyogakgun_total = defaultdict(int)
kyogakgun_sudo = defaultdict(int)
kyogakgun_by_sido = defaultdict(lambda: defaultdict(int))
for subj, qty in subject_total.items():
    kg = kyogakgun_map.get(subj, '기타')
    kyogakgun_total[kg] += qty
for subj, qty in subject_sudogwon.items():
    kg = kyogakgun_map.get(subj, '기타')
    kyogakgun_sudo[kg] += qty
for sido in sido_covered:
    for subj, qty in subject_by_sido[sido].items():
        kg = kyogakgun_map.get(subj, '기타')
        kyogakgun_by_sido[sido][kg] += qty

kyogakgun_nationwide = [
    {'kg': kg, 'qty': qty, 'pct': safe_pct(qty, total_all_students)}
    for kg, qty in sorted(kyogakgun_total.items(), key=lambda x: -x[1])
]
kyogakgun_sudogwon = [
    {'kg': kg, 'qty': qty, 'pct': safe_pct(qty, total_sudo_students)}
    for kg, qty in sorted(kyogakgun_sudo.items(), key=lambda x: -x[1])
]
kyogakgun_regional = [
    {
        'sido': sido,
        'kgs': [{'kg': kg, 'qty': qty, 'pct': safe_pct(qty, (sido_students.get(sido, {}).get('g1', 0)+sido_students.get(sido, {}).get('g2', 0)+sido_students.get(sido, {}).get('g3', 0))/3.0)}
                for kg, qty in sorted(kyogakgun_by_sido[sido].items(), key=lambda x: -x[1])]
    }
    for sido in sorted(sido_covered)
]

# Region (구) subject shares - 고등학교
region_shares = []
for region in sorted(region_covered):
    rcohort = region_students.get(region, 0)
    subjects_in_region = sorted(subject_by_region[region].items(), key=lambda x: -x[1])
    subj_list = []
    for subj, qty in subjects_in_region:
        kg = kyogakgun_map.get(subj, '기타')
        pct = safe_pct(qty, rcohort) if rcohort > 0 else 0.0
        subj_list.append({'subject': subj, 'kyogakgun': kg, 'qty': qty, 'pct': pct})
    region_shares.append({
        'region': region,
        'cohort_students': round(rcohort, 1),
        'subjects': subj_list,
    })

# Region (구) subject shares - 중학교
ms_region_shares = []
for region in sorted(ms_region_covered):
    rcohort = ms_region_students.get(region, 0)
    subjects_in_region = sorted(ms_subject_by_region[region].items(), key=lambda x: -x[1])
    subj_list = []
    for subj, qty in subjects_in_region:
        kg = ms_kyogakgun_map.get(subj, '기타')
        pct = safe_pct(qty, rcohort) if rcohort > 0 else 0.0
        subj_list.append({'subject': subj, 'kyogakgun': kg, 'qty': qty, 'pct': pct})
    ms_region_shares.append({
        'region': region,
        'cohort_students': round(rcohort, 1),
        'subjects': subj_list,
    })

# 중학교 전체/수도권/지방 shares
# 전체 중학교 cohort: ms_region_students 합산
ms_nw_cohort   = sum(ms_region_students.values()) or 1
ms_su_cohort   = sum(v for r, v in ms_region_students.items()
                     if any(r.startswith(p) for p in ['서울특별시', '인천광역시', '경기도'])) or 1
ms_jb_cohort   = ms_nw_cohort - ms_su_cohort or 1

ms_nationwide_shares = [
    {'subject': s, 'kyogakgun': ms_kyogakgun_map.get(s, ''), 'qty': q,
     'pct': safe_pct(q, ms_nw_cohort)}
    for s, q in sorted(ms_subject_total.items(), key=lambda x: -x[1])
]
ms_sudogwon_shares = [
    {'subject': s, 'kyogakgun': ms_kyogakgun_map.get(s, ''), 'qty': q,
     'pct': safe_pct(q, ms_su_cohort)}
    for s, q in sorted(ms_subject_sudogwon.items(), key=lambda x: -x[1])
]
ms_jibang_shares = [
    {'subject': s, 'kyogakgun': ms_kyogakgun_map.get(s, ''), 'qty': q,
     'pct': safe_pct(q, ms_jb_cohort)}
    for s, q in sorted(ms_subject_jibang.items(), key=lambda x: -x[1])
]
print(f"MS nationwide shares: {len(ms_nationwide_shares)} subjects")
print(f"MS region shares: {len(ms_region_shares)} regions")

# ------------------------------------------------------------------
# 3b. Jibang (지방 = 비수도권) shares
# ------------------------------------------------------------------
jibang_todos = {s for s in sido_covered if s not in SUDOGWON}
jibang_students_g1 = sum(sido_students.get(s, {}).get('g1', 0) for s in jibang_todos)
jibang_students_g2 = sum(sido_students.get(s, {}).get('g2', 0) for s in jibang_todos)
jibang_students_g3 = sum(sido_students.get(s, {}).get('g3', 0) for s in jibang_todos)
jibang_cohort = (jibang_students_g1 + jibang_students_g2 + jibang_students_g3) / 3.0
jibang_shares = [
    {
        'subject': subj,
        'kyogakgun': kyogakgun_map.get(subj, '기타'),
        'qty': qty,
        'pct': safe_pct(qty, jibang_cohort),
    }
    for subj, qty in sorted(subject_jibang.items(), key=lambda x: -x[1])
]
print(f"Jibang cohort: {jibang_cohort:,.0f}, subjects: {len(jibang_shares)}")

# ------------------------------------------------------------------
# 3c. Region detail: school-level order + grade
# ------------------------------------------------------------------
def school_grade_tier(students):
    if students >= 800: return 'S'
    if students >= 500: return 'A'
    if students >= 250: return 'B'
    return 'C'

region_detail = {}
# 담당자 구 목록에 있는 지역만 생성 (데이터 크기 절약)
all_staff_regions = set()
for regs in staff_regions_map.values():
    all_staff_regions.update(regs)

for region in all_staff_regions:
    schools_meta = region_school_meta.get(region, [])
    if not schools_meta:
        continue
    school_list = []
    for sch in schools_meta:
        code = sch['code']
        orders = subject_by_school.get(code, {})
        # cohort: school students / 3
        sch_cohort = sch['students'] / 3.0 if sch['students'] else 0
        ordered_subjects = [
            {
                'subject': subj,
                'qty': qty,
                'kyogakgun': kyogakgun_map.get(subj, '기타'),
                'pct': safe_pct(qty, sch_cohort) if sch_cohort > 0 else 0.0,
            }
            for subj, qty in sorted(orders.items(), key=lambda x: -x[1])
        ]
        school_list.append({
            'name': sch['name'],
            'code': code,
            'students': sch['students'],
            'type': sch['type'],
            'founded': sch['founded'],
            'grade': school_grade_tier(sch['students']),
            'has_order': len(orders) > 0,
            'ordered_subjects': ordered_subjects,
        })
    school_list.sort(key=lambda x: (-x['students']))
    region_detail[region] = school_list

print(f"Region detail computed for {len(region_detail)} regions")

# Build output JSON
# ------------------------------------------------------------------
out = {
    'meta': {
        'year': 2026,
        'description': '2026학년도 고등+중학교 교과서 채택 점유율',
        'source': '21-26주문데이터.xlsx + 학교별 담당총판 및 본사담당자 정보.xlsx',
    },
    'students': {
        'nationwide': total_students_all,
        'sudogwon': sudogwon_students,
        'by_sido': {sido: v for sido, v in sido_students.items()},
    },
    'nationwide_shares': nationwide_shares,
    'sudogwon_shares': sudogwon_shares,
    'jibang_shares': jibang_shares,
    'regional': regional,
    'region_shares': region_shares,
    'region_detail': region_detail,
    'staff_regions': staff_regions_map,
    'kyogakgun': {
        'nationwide': kyogakgun_nationwide,
        'sudogwon': kyogakgun_sudogwon,
        'regional': kyogakgun_regional,
    },
    # 중학교 데이터
    'ms': {
        'nationwide_shares': ms_nationwide_shares,
        'sudogwon_shares': ms_sudogwon_shares,
        'jibang_shares': ms_jibang_shares,
        'region_shares': ms_region_shares,
        'cohort': {
            'nationwide': round(ms_nw_cohort, 1),
            'sudogwon': round(ms_su_cohort, 1),
            'jibang': round(ms_jb_cohort, 1),
        },
    },
}

with open(OUT_FILE, 'w', encoding='utf-8') as f:
    json.dump(out, f, ensure_ascii=False, indent=2)

print(f"\nSaved to: {OUT_FILE}")
print(f"  HS Nationwide subjects: {len(nationwide_shares)}")
print(f"  HS Regions: {len(regional)}")
print(f"  HS District regions: {len(region_shares)}")
print(f"  MS Nationwide subjects: {len(ms_nationwide_shares)}")
print(f"  MS District regions: {len(ms_region_shares)}")
